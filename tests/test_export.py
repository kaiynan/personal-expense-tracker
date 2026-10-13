"""Pengujian unit untuk fitur V2 Export CSV, Excel (.xlsx), dan PDF (.pdf)."""

from __future__ import annotations

import csv
from datetime import date

import pytest
from openpyxl import load_workbook

import config
from database.engine import create_session_factory, session_scope
from database.seed import ensure_default_user, seed_defaults
from models import TYPE_EXPENSE, TYPE_INCOME
from reports import (
    CSV_HEADERS,
    export_csv,
    export_excel,
    export_pdf,
    get_export_data,
)
from services import adjustments, balances, categories, opening_balances, transactions, transfers
from utils.errors import NotFoundError, ValidationError


@pytest.fixture
def exp_ctx(engine, tmp_path, monkeypatch):
    """Factory session dan user_id dengan REPORTS_DIR diarahkan ke tmp_path."""
    monkeypatch.setattr(config, "REPORTS_DIR", tmp_path / "reports_output")
    factory = create_session_factory(engine)
    with session_scope(factory) as session:
        user = ensure_default_user(session)
        seed_defaults(session, user.id)
        u_id = user.id

        methods = balances.get_method_balances(session, u_id)
        cash_id = next(m.method_id for m in methods if m.name == "Cash")
        bank_id = next(m.method_id for m in methods if m.name == "Bank")

        inc_cats = categories.list_categories(session, u_id, type_=TYPE_INCOME)
        exp_cats = categories.list_categories(session, u_id, type_=TYPE_EXPENSE)
        salary_id = next(c.id for c in inc_cats if c.name == "Gaji")
        food_id = next(c.id for c in exp_cats if c.name == "Makanan")
        transport_id = next(c.id for c in exp_cats if c.name == "Transportasi")

        opening_balances.set_opening_balance(
            session, u_id, bank_id, 2_000_000, effective_date="01-10-2026"
        )
        opening_balances.set_opening_balance(
            session, u_id, cash_id, 500_000, effective_date="01-10-2026"
        )

        # 1. Pemasukan 3.000.000
        transactions.add_income(
            session, u_id, 3_000_000, salary_id, bank_id, on="05-10-2026", note="Gaji Oktober"
        )
        # 2. Pengeluaran Makanan 150.000
        transactions.add_expense(
            session, u_id, 150_000, food_id, cash_id, on="06-10-2026", note="Makan siang"
        )
        # 3. Pengeluaran Transportasi 50.000
        transactions.add_expense(
            session, u_id, 50_000, transport_id, cash_id, on="07-10-2026", note="Bensin motor"
        )
        # 4. Transfer Bank -> Cash 400.000 dengan biaya admin 2.500
        transfers.add_transfer(
            session,
            u_id,
            400_000,
            bank_id,
            cash_id,
            on="08-10-2026",
            admin_fee=2_500,
            note="Tarik tunai ATM",
        )
        # 5. Adjustment pada Cash (saldo saat ini 700k -> koreksi jadi 690k, selisih -10k)
        adjustments.adjust_balance(
            session, u_id, cash_id, 690_000, on="09-10-2026", note="Selisih uang receh"
        )
        # 6. Transaksi soft-deleted (tidak boleh masuk ekspor)
        del_res = transactions.add_expense(
            session, u_id, 999_000, food_id, cash_id, on="10-10-2026", note="Salah input"
        )
        transactions.delete_transaction(session, u_id, del_res.transaction.id)

    return factory, u_id


def test_get_export_data_domain_rules(exp_ctx):
    """Ringkasan hanya menghitung income & expense aktif (tanpa transfer/adjustment/deleted)."""
    factory, user_id = exp_ctx
    with session_scope(factory) as session:
        payload = get_export_data(session, user_id)

    # Total income = 3.000.000
    assert payload.total_income == 3_000_000
    # Total expense = 150.000 + 50.000 + 2.500 (biaya admin) = 202.500
    assert payload.total_expense == 202_500
    assert payload.net == 3_000_000 - 202_500
    assert payload.income_expense_count == 4

    # Jumlah baris transaksi aktif = 6 (1 inc, 2 exp, 1 transfer, 1 biaya admin, 1 adjustment)
    assert len(payload.transactions) == 6
    notes = [r.note for r in payload.transactions]
    assert "Salah input" not in notes


def test_export_csv_creates_valid_file(exp_ctx):
    """File CSV berisi header standar, format tanggal DD-MM-YYYY, dan nominal integer murni."""
    factory, user_id = exp_ctx
    with session_scope(factory) as session:
        csv_path = export_csv(session, user_id)

    assert csv_path.exists()
    assert csv_path.suffix == ".csv"

    with csv_path.open("r", encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))

    assert rows[0] == CSV_HEADERS
    # 1 baris header + 6 baris transaksi aktif
    assert len(rows) == 7

    # Cek baris terbaru (adjustment pada 09-10-2026)
    first_data = rows[1]
    assert first_data[1] == "09-10-2026"
    assert first_data[2] == "adjustment"
    assert first_data[6] == "-10000"

    # Pastikan baris transfer memiliki metode tujuan
    transfer_rows = [r for r in rows[1:] if r[2] == "transfer"]
    assert len(transfer_rows) == 1
    assert transfer_rows[0][4] == "Bank"
    assert transfer_rows[0][5] == "Cash"
    assert transfer_rows[0][6] == "400000"


def test_export_excel_creates_multi_sheet_workbook(exp_ctx):
    """File Excel (.xlsx) memiliki 3 sheet dengan ringkasan laporan dan nominal bertipe int."""
    factory, user_id = exp_ctx
    with session_scope(factory) as session:
        xlsx_path = export_excel(session, user_id)

    assert xlsx_path.exists()
    assert xlsx_path.suffix == ".xlsx"

    wb = load_workbook(xlsx_path)
    assert wb.sheetnames == ["Ringkasan Laporan", "Daftar Transaksi", "Pengeluaran Terbesar"]

    ws_sum = wb["Ringkasan Laporan"]
    # Baris 7 = Total Pemasukan, Baris 8 = Total Pengeluaran, Baris 9 = Selisih (Net)
    assert ws_sum["A7"].value == "Total Pemasukan"
    assert ws_sum["B7"].value == 3_000_000
    assert isinstance(ws_sum["B7"].value, int)
    assert ws_sum["A8"].value == "Total Pengeluaran"
    assert ws_sum["B8"].value == 202_500
    assert ws_sum["B9"].value == 3_000_000 - 202_500

    ws_tx = wb["Daftar Transaksi"]
    # 1 header + 6 transaksi
    assert ws_tx.max_row == 7
    assert isinstance(ws_tx.cell(row=2, column=7).value, int)

    ws_top = wb["Pengeluaran Terbesar"]
    # 1 header + 3 pengeluaran (Makanan 150k, Transportasi 50k, Biaya Admin 2.5k)
    assert ws_top.max_row == 4
    assert ws_top.cell(row=2, column=4).value == 150_000
    wb.close()


def test_export_pdf_creates_valid_pdf_and_handles_unicode(exp_ctx):
    """File PDF (.pdf) berhasil dibuat, diawali header %PDF-, dan tahan karakter emoji/Unicode."""
    factory, user_id = exp_ctx
    with session_scope(factory) as session:
        exp_cat = categories.list_categories(session, user_id, type_=TYPE_EXPENSE)[0]
        pm = balances.get_method_balances(session, user_id)[0]
        transactions.add_expense(
            session,
            user_id,
            25_000,
            exp_cat.id,
            pm.method_id,
            on="11-10-2026",
            note="Kopi susu ☕🔥 mantap",
        )
        pdf_path = export_pdf(session, user_id)

    assert pdf_path.exists()
    assert pdf_path.suffix == ".pdf"
    raw = pdf_path.read_bytes()
    assert raw.startswith(b"%PDF-")
    assert len(raw) > 500


def test_export_filters_and_validation(exp_ctx, tmp_path):
    """Filter rentang tanggal, jenis transaksi, dan validasi input tanggal/path bekerja."""
    factory, user_id = exp_ctx
    with session_scope(factory) as session:
        # Filter rentang 06-10-2026 s.d. 07-10-2026
        payload = get_export_data(session, user_id, start=date(2026, 10, 6), end="07-10-2026")
        assert len(payload.transactions) == 2
        assert payload.total_expense == 200_000
        assert payload.total_income == 0

        # Validasi start > end
        with pytest.raises(ValidationError):
            export_csv(session, user_id, start="10-10-2026", end="01-10-2026")

        # Validasi file_path kosong
        with pytest.raises(ValidationError):
            export_csv(session, user_id, file_path="   ")

        # Validasi user tidak ditemukan
        with pytest.raises(NotFoundError):
            export_pdf(session, 999999)

        # Custom file path tanpa ekstensi otomatis diberi ekstensi yang tepat
        custom_xlsx = export_excel(session, user_id, file_path=tmp_path / "laporan_oktober")
        assert custom_xlsx.suffix == ".xlsx"
        assert custom_xlsx.exists()
