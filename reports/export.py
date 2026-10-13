"""Modul export data dan laporan keuangan ke format CSV, Excel (.xlsx), dan PDF (.pdf).

Mematuhi aturan domain:
- Semua nominal adalah integer rupiah (tanpa float).
- Hanya mengekspor transaksi berstatus `active`.
- Transfer dan adjustment tercatat di daftar transaksi, tetapi TIDAK dihitung
  sebagai pemasukan maupun pengeluaran pada ringkasan laporan.
- Setiap query dan operasi diisolasi berdasarkan `user_id`.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

import config
from database import queries
from models import TYPE_EXPENSE, TYPE_INCOME
from reports.financial import CategoryExpenseItem, TopExpenseItem
from services import balances, transactions
from services.balances import MethodBalance
from services.common import require_user
from utils.dates import format_display
from utils.errors import ValidationError
from utils.formatting import format_rupiah
from utils.validators import parse_date

CSV_HEADERS: list[str] = [
    "ID",
    "Tanggal",
    "Jenis",
    "Kategori",
    "Metode Pembayaran",
    "Metode Tujuan",
    "Nominal (Rp)",
    "Nominal Terformat",
    "Keterangan",
]


def sanitize_cell_value(val: Any) -> Any:
    """Netralkan nilai sel yang diawali =, +, -, atau @ pada ekspor (formula injection)."""
    if isinstance(val, str) and val and val[0] in ("=", "+", "-", "@"):
        return f"'{val}"
    return val


@dataclass(frozen=True)
class ExportTransactionRow:
    """Satu baris transaksi siap ekspor."""

    id: int
    date: date
    date_display: str
    type: str
    category_name: str
    method_name: str
    to_method_name: str
    amount: int
    amount_formatted: str
    note: str


@dataclass(frozen=True)
class ExportPayload:
    """Kumpulan data lengkap laporan dan daftar transaksi untuk ekspor."""

    user_id: int
    user_name: str
    period_label: str
    start_date: date | None
    end_date: date | None
    total_balance: int
    method_balances: list[MethodBalance]
    total_income: int
    total_expense: int
    net: int
    income_expense_count: int
    category_expenses: list[CategoryExpenseItem]
    top_expenses: list[TopExpenseItem]
    transactions: list[ExportTransactionRow]


def _resolve_output_path(
    user_id: int,
    ext: str,
    file_path: str | Path | None = None,
    start_date: date | None = None,
    end_date: date | None = None,
) -> Path:
    """Tentukan path file output di `config.REPORTS_DIR` atau path kustom."""
    config.ensure_directories()
    clean_ext = ext if ext.startswith(".") else f".{ext}"

    if file_path is not None:
        raw_str = str(file_path).strip()
        if not raw_str:
            raise ValidationError("Path file tujuan ekspor tidak boleh kosong.")
        candidate = Path(raw_str)
        if candidate.suffix.lower() != clean_ext.lower():
            candidate = candidate.with_suffix(clean_ext)
        if not candidate.is_absolute() and len(candidate.parts) == 1:
            candidate = config.REPORTS_DIR / candidate
        candidate.parent.mkdir(parents=True, exist_ok=True)
        return candidate

    ts = datetime.now(UTC).strftime("%Y%m%d_%H%M%S_%f")
    if start_date and end_date:
        period_slug = f"{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}"
    elif start_date:
        period_slug = f"dari_{start_date.strftime('%Y%m%d')}"
    elif end_date:
        period_slug = f"sampai_{end_date.strftime('%Y%m%d')}"
    else:
        period_slug = "semua"

    filename = f"export_user{user_id}_{period_slug}_{ts}{clean_ext}"
    return config.REPORTS_DIR / filename


def _sanitize_pdf_text(text: str) -> str:
    """Bersihkan karakter non-Latin-1 agar aman dirender oleh font bawaan Helvetica di FPDF."""
    return text.encode("latin-1", errors="replace").decode("latin-1")


def get_export_data(
    session: Session,
    user_id: int,
    *,
    start: date | str | None = None,
    end: date | str | None = None,
    type_: str | None = None,
    category_id: int | None = None,
    method_id: int | None = None,
    keyword: str | None = None,
) -> ExportPayload:
    """Kumpulkan dan susun data transaksi serta ringkasan laporan untuk ekspor."""
    user = require_user(session, user_id)

    start_date = parse_date(start, "Tanggal awal") if start is not None else None
    end_date = parse_date(end, "Tanggal akhir") if end is not None else None
    if start_date and end_date and start_date > end_date:
        raise ValidationError("Tanggal awal tidak boleh setelah tanggal akhir.")

    tx_list = transactions.list_transactions(
        session,
        user_id,
        type_=type_,
        category_id=category_id,
        method_id=method_id,
        start=start_date,
        end=end_date,
        keyword=keyword,
    )

    cat_map = {
        c.id: c.name
        for c in queries.list_categories(
            session, user_id, include_archived=True, include_system=True
        )
    }
    method_map = {m.id: m.name for m in queries.list_methods(session, user_id)}

    rows: list[ExportTransactionRow] = []
    total_income = 0
    total_expense = 0
    income_expense_count = 0
    cat_totals: dict[int, int] = {}

    for tx in tx_list:
        cname = cat_map.get(tx.category_id, "-") if tx.category_id is not None else "-"
        mname = method_map.get(tx.method_id, "-")
        to_mname = method_map.get(tx.to_method_id, "-") if tx.to_method_id is not None else "-"

        rows.append(
            ExportTransactionRow(
                id=tx.id,
                date=tx.date,
                date_display=format_display(tx.date),
                type=tx.type,
                category_name=cname,
                method_name=mname,
                to_method_name=to_mname,
                amount=int(tx.amount),
                amount_formatted=format_rupiah(int(tx.amount)),
                note=tx.note or "",
            )
        )

        # Aturan domain #3 & #9: hanya income dan expense yang masuk hitungan laporan
        if tx.type == TYPE_INCOME:
            total_income += int(tx.amount)
            income_expense_count += 1
        elif tx.type == TYPE_EXPENSE:
            amt = int(tx.amount)
            total_expense += amt
            income_expense_count += 1
            if tx.category_id is not None:
                cat_totals[tx.category_id] = cat_totals.get(tx.category_id, 0) + amt

    category_expenses = [
        CategoryExpenseItem(
            category_id=cid,
            category_name=cat_map.get(cid, "-"),
            amount=amt,
            percentage=round((amt / total_expense) * 100, 1) if total_expense > 0 else 0.0,
        )
        for cid, amt in sorted(cat_totals.items(), key=lambda item: (-item[1], item[0]))
    ]

    expense_txs = [tx for tx in tx_list if tx.type == TYPE_EXPENSE]
    expense_txs.sort(key=lambda t: (-int(t.amount), -t.date.toordinal(), -t.id))
    top_expenses = [
        TopExpenseItem(
            id=tx.id,
            date=tx.date,
            category_name=cat_map.get(tx.category_id, "-") if tx.category_id is not None else "-",
            amount=int(tx.amount),
            note=tx.note,
        )
        for tx in expense_txs[:5]
    ]

    if start_date and end_date:
        if start_date == end_date:
            period_label = format_display(start_date)
        else:
            period_label = f"{format_display(start_date)} s.d. {format_display(end_date)}"
    elif start_date:
        period_label = f"Mulai {format_display(start_date)}"
    elif end_date:
        period_label = f"Sampai {format_display(end_date)}"
    else:
        period_label = "Semua Periode"

    method_bals = balances.get_method_balances(session, user_id)
    total_bal = balances.get_total_balance(session, user_id)

    return ExportPayload(
        user_id=user.id,
        user_name=user.name,
        period_label=period_label,
        start_date=start_date,
        end_date=end_date,
        total_balance=total_bal,
        method_balances=method_bals,
        total_income=total_income,
        total_expense=total_expense,
        net=total_income - total_expense,
        income_expense_count=income_expense_count,
        category_expenses=category_expenses,
        top_expenses=top_expenses,
        transactions=rows,
    )


def export_csv(
    session: Session,
    user_id: int,
    *,
    start: date | str | None = None,
    end: date | str | None = None,
    type_: str | None = None,
    category_id: int | None = None,
    method_id: int | None = None,
    keyword: str | None = None,
    file_path: str | Path | None = None,
) -> Path:
    """Ekspor daftar transaksi pengguna ke file CSV (menggunakan modul standar `csv`)."""
    payload = get_export_data(
        session,
        user_id,
        start=start,
        end=end,
        type_=type_,
        category_id=category_id,
        method_id=method_id,
        keyword=keyword,
    )
    target_path = _resolve_output_path(
        user_id,
        ".csv",
        file_path=file_path,
        start_date=payload.start_date,
        end_date=payload.end_date,
    )

    with target_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(CSV_HEADERS)
        for row in payload.transactions:
            writer.writerow(
                [
                    row.id,
                    sanitize_cell_value(row.date_display),
                    sanitize_cell_value(row.type),
                    sanitize_cell_value(row.category_name),
                    sanitize_cell_value(row.method_name),
                    sanitize_cell_value(row.to_method_name),
                    row.amount,
                    sanitize_cell_value(row.amount_formatted),
                    sanitize_cell_value(row.note),
                ]
            )

    return target_path


def export_excel(
    session: Session,
    user_id: int,
    *,
    start: date | str | None = None,
    end: date | str | None = None,
    type_: str | None = None,
    category_id: int | None = None,
    method_id: int | None = None,
    keyword: str | None = None,
    file_path: str | Path | None = None,
) -> Path:
    """Ekspor ringkasan laporan dan daftar transaksi pengguna ke file Excel (.xlsx)."""
    payload = get_export_data(
        session,
        user_id,
        start=start,
        end=end,
        type_=type_,
        category_id=category_id,
        method_id=method_id,
        keyword=keyword,
    )
    target_path = _resolve_output_path(
        user_id,
        ".xlsx",
        file_path=file_path,
        start_date=payload.start_date,
        end_date=payload.end_date,
    )

    wb = Workbook()

    # Styling dasar
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    section_font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2B5C8F", end_color="2B5C8F", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    thin_side = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    # --- Sheet 1: Ringkasan Laporan ---
    ws_summary = wb.active
    ws_summary.title = "Ringkasan Laporan"

    ws_summary["A1"] = "LAPORAN KEUANGAN PERSONAL EXPENSE TRACKER"
    ws_summary["A1"].font = title_font
    ws_summary["A2"] = f"Pengguna: {payload.user_name}"
    ws_summary["A3"] = f"Periode: {payload.period_label}"

    # Ringkasan Arus Kas
    ws_summary["A5"] = "RINGKASAN ARUS KAS"
    ws_summary["A5"].font = section_font

    summary_rows = [
        ("Total Pemasukan", payload.total_income, format_rupiah(payload.total_income)),
        ("Total Pengeluaran", payload.total_expense, format_rupiah(payload.total_expense)),
        ("Selisih (Net)", payload.net, format_rupiah(payload.net)),
        (
            "Jumlah Transaksi (Income/Expense)",
            payload.income_expense_count,
            f"{payload.income_expense_count} transaksi",
        ),
        (
            "Total Baris Transaksi Diekspor",
            len(payload.transactions),
            f"{len(payload.transactions)} baris",
        ),
    ]
    ws_summary.append(["Komponen", "Nominal (Rp)", "Tampilan"])
    for col_idx in range(1, 4):
        cell = ws_summary.cell(row=6, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for label, val_int, val_str in summary_rows:
        ws_summary.append([label, val_int, val_str])
        r_idx = ws_summary.max_row
        ws_summary.cell(row=r_idx, column=2).number_format = "#,##0"
        for c_idx in range(1, 4):
            ws_summary.cell(row=r_idx, column=c_idx).border = cell_border

    # Saldo Saat Ini per Metode
    start_bal_row = ws_summary.max_row + 2
    bal_title = "SALDO REKENING / METODE PEMBAYARAN"
    ws_summary.cell(row=start_bal_row, column=1, value=bal_title).font = section_font
    hdr_bal_row = start_bal_row + 1
    ws_summary.cell(row=hdr_bal_row, column=1, value="Metode Pembayaran")
    ws_summary.cell(row=hdr_bal_row, column=2, value="Saldo (Rp)")
    ws_summary.cell(row=hdr_bal_row, column=3, value="Tampilan")
    for c_idx in range(1, 4):
        c = ws_summary.cell(row=hdr_bal_row, column=c_idx)
        c.font = header_font
        c.fill = header_fill

    for mb in payload.method_balances:
        if mb.is_archived and mb.balance == 0:
            continue
        ws_summary.append([mb.name, mb.balance, format_rupiah(mb.balance)])
        r_idx = ws_summary.max_row
        ws_summary.cell(row=r_idx, column=2).number_format = "#,##0"
        for c_idx in range(1, 4):
            ws_summary.cell(row=r_idx, column=c_idx).border = cell_border

    ws_summary.append(["TOTAL SALDO", payload.total_balance, format_rupiah(payload.total_balance)])
    tot_r = ws_summary.max_row
    ws_summary.cell(row=tot_r, column=2).number_format = "#,##0"
    for c_idx in range(1, 4):
        cell = ws_summary.cell(row=tot_r, column=c_idx)
        cell.font = bold_font
        cell.border = cell_border

    # Pengeluaran per Kategori
    start_cat_row = ws_summary.max_row + 2
    cat_title = "PENGELUARAN PER KATEGORI"
    ws_summary.cell(row=start_cat_row, column=1, value=cat_title).font = section_font
    hdr_cat_row = start_cat_row + 1
    cat_headers = ["Kategori", "Nominal (Rp)", "Tampilan", "Persentase (%)"]
    for c_idx, h_text in enumerate(cat_headers, start=1):
        c = ws_summary.cell(row=hdr_cat_row, column=c_idx, value=h_text)
        c.font = header_font
        c.fill = header_fill

    if payload.category_expenses:
        for item in payload.category_expenses:
            ws_summary.append(
                [
                    item.category_name,
                    item.amount,
                    format_rupiah(item.amount),
                    f"{item.percentage:.1f}%",
                ]
            )
            r_idx = ws_summary.max_row
            ws_summary.cell(row=r_idx, column=2).number_format = "#,##0"
            for c_idx in range(1, 5):
                ws_summary.cell(row=r_idx, column=c_idx).border = cell_border
    else:
        ws_summary.append(["(Tidak ada pengeluaran pada periode ini)", 0, "Rp0", "0.0%"])

    # --- Sheet 2: Daftar Transaksi ---
    ws_tx = wb.create_sheet(title="Daftar Transaksi")
    ws_tx.append(CSV_HEADERS)
    for col_idx in range(1, len(CSV_HEADERS) + 1):
        c = ws_tx.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    for row in payload.transactions:
        ws_tx.append(
            [
                row.id,
                sanitize_cell_value(row.date_display),
                sanitize_cell_value(row.type),
                sanitize_cell_value(row.category_name),
                sanitize_cell_value(row.method_name),
                sanitize_cell_value(row.to_method_name),
                row.amount,
                sanitize_cell_value(row.amount_formatted),
                sanitize_cell_value(row.note),
            ]
        )
        r_idx = ws_tx.max_row
        ws_tx.cell(row=r_idx, column=7).number_format = "#,##0"
        for c_idx in range(1, len(CSV_HEADERS) + 1):
            ws_tx.cell(row=r_idx, column=c_idx).border = cell_border

    # --- Sheet 3: Pengeluaran Terbesar ---
    ws_top = wb.create_sheet(title="Pengeluaran Terbesar")
    top_headers = ["ID", "Tanggal", "Kategori", "Nominal (Rp)", "Nominal Terformat", "Keterangan"]
    ws_top.append(top_headers)
    for col_idx in range(1, len(top_headers) + 1):
        c = ws_top.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill

    for te in payload.top_expenses:
        ws_top.append(
            [
                te.id,
                sanitize_cell_value(format_display(te.date)),
                sanitize_cell_value(te.category_name),
                te.amount,
                sanitize_cell_value(format_rupiah(te.amount)),
                sanitize_cell_value(te.note or ""),
            ]
        )
        r_idx = ws_top.max_row
        ws_top.cell(row=r_idx, column=4).number_format = "#,##0"
        for c_idx in range(1, len(top_headers) + 1):
            ws_top.cell(row=r_idx, column=c_idx).border = cell_border

    # Auto-fit lebar kolom pada semua sheet
    for ws in (ws_summary, ws_tx, ws_top):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    wb.save(target_path)
    wb.close()
    return target_path


def export_pdf(
    session: Session,
    user_id: int,
    *,
    start: date | str | None = None,
    end: date | str | None = None,
    type_: str | None = None,
    category_id: int | None = None,
    method_id: int | None = None,
    keyword: str | None = None,
    file_path: str | Path | None = None,
) -> Path:
    """Ekspor laporan keuangan dan daftar transaksi pengguna ke file PDF (.pdf)."""
    payload = get_export_data(
        session,
        user_id,
        start=start,
        end=end,
        type_=type_,
        category_id=category_id,
        method_id=method_id,
        keyword=keyword,
    )
    target_path = _resolve_output_path(
        user_id,
        ".pdf",
        file_path=file_path,
        start_date=payload.start_date,
        end_date=payload.end_date,
    )

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Judul Dokumen
    pdf.set_font("Helvetica", style="B", size=15)
    pdf.cell(
        0,
        9,
        _sanitize_pdf_text("LAPORAN KEUANGAN & TRANSAKSI"),
        new_x="LMARGIN",
        new_y="NEXT",
        align="C",
    )

    pdf.set_font("Helvetica", size=10)
    pdf.cell(
        0,
        6,
        _sanitize_pdf_text(f"Pengguna: {payload.user_name}  |  Periode: {payload.period_label}"),
        new_x="LMARGIN",
        new_y="NEXT",
        align="C",
    )
    pdf.ln(4)

    # 1. Ringkasan Arus Kas & Saldo
    pdf.set_font("Helvetica", style="B", size=11)
    pdf.cell(0, 7, "1. Ringkasan Arus Kas & Saldo", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", size=10)
    net_sign = "+" if payload.net > 0 else ""
    summary_lines = [
        ("Total Pemasukan", format_rupiah(payload.total_income)),
        ("Total Pengeluaran", format_rupiah(payload.total_expense)),
        ("Selisih (Arus Kas Bersih)", f"{net_sign}{format_rupiah(payload.net)}"),
        ("Jumlah Transaksi (Pemasukan/Pengeluaran)", f"{payload.income_expense_count} transaksi"),
        ("Total Saldo Saat Ini", format_rupiah(payload.total_balance)),
    ]
    for label, val in summary_lines:
        pdf.cell(85, 6, _sanitize_pdf_text(label), border=1)
        pdf.cell(60, 6, _sanitize_pdf_text(val), border=1, new_x="LMARGIN", new_y="NEXT")

    pdf.ln(4)

    # 2. Pengeluaran per Kategori
    pdf.set_font("Helvetica", style="B", size=11)
    pdf.cell(0, 7, "2. Pengeluaran per Kategori", new_x="LMARGIN", new_y="NEXT")

    if payload.category_expenses:
        pdf.set_font("Helvetica", style="B", size=9)
        pdf.cell(75, 6, "Kategori", border=1)
        pdf.cell(45, 6, "Nominal", border=1, align="R")
        pdf.cell(25, 6, "Porsi (%)", border=1, new_x="LMARGIN", new_y="NEXT", align="R")

        pdf.set_font("Helvetica", size=9)
        for cat in payload.category_expenses:
            pdf.cell(75, 6, _sanitize_pdf_text(cat.category_name[:38]), border=1)
            pdf.cell(45, 6, _sanitize_pdf_text(format_rupiah(cat.amount)), border=1, align="R")
            pdf.cell(
                25,
                6,
                f"{cat.percentage:.1f}%",
                border=1,
                new_x="LMARGIN",
                new_y="NEXT",
                align="R",
            )
    else:
        pdf.set_font("Helvetica", style="I", size=9)
        pdf.cell(
            0,
            6,
            "Tidak ada data pengeluaran pada periode ini.",
            new_x="LMARGIN",
            new_y="NEXT",
        )

    pdf.ln(4)

    # 3. Daftar Transaksi
    pdf.set_font("Helvetica", style="B", size=11)
    pdf.cell(
        0,
        7,
        f"3. Daftar Transaksi ({len(payload.transactions)} baris)",
        new_x="LMARGIN",
        new_y="NEXT",
    )

    if payload.transactions:
        pdf.set_font("Helvetica", style="B", size=8)
        col_widths = (12, 22, 20, 32, 28, 32, 44)
        col_headers = ("ID", "Tanggal", "Jenis", "Kategori", "Metode", "Nominal", "Keterangan")
        for w, h in zip(col_widths, col_headers, strict=True):
            pdf.cell(w, 6, h, border=1, align="C")
        pdf.ln()

        pdf.set_font("Helvetica", size=8)
        for row in payload.transactions:
            method_label = (
                f"{row.method_name}->{row.to_method_name}"
                if row.to_method_name != "-"
                else row.method_name
            )
            note_short = row.note[:26] + "..." if len(row.note) > 29 else row.note

            pdf.cell(col_widths[0], 6, f"#{row.id}", border=1, align="C")
            pdf.cell(col_widths[1], 6, row.date_display, border=1, align="C")
            pdf.cell(col_widths[2], 6, _sanitize_pdf_text(row.type), border=1, align="C")
            pdf.cell(
                col_widths[3], 6, _sanitize_pdf_text(row.category_name[:18]), border=1, align="L"
            )
            pdf.cell(col_widths[4], 6, _sanitize_pdf_text(method_label[:16]), border=1, align="L")
            pdf.cell(
                col_widths[5], 6, _sanitize_pdf_text(row.amount_formatted), border=1, align="R"
            )
            pdf.cell(
                col_widths[6],
                6,
                _sanitize_pdf_text(note_short),
                border=1,
                new_x="LMARGIN",
                new_y="NEXT",
                align="L",
            )
    else:
        pdf.set_font("Helvetica", style="I", size=9)
        pdf.cell(0, 6, "Tidak ada transaksi yang ditemukan.", new_x="LMARGIN", new_y="NEXT")

    pdf.output(str(target_path))
    return target_path
